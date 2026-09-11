# -*- coding: utf-8 -*-
"""Parse Feishu Sprint export xlsx bundle into normalized sprint summary dataset."""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ITEM_DONE = {"已验收", "待验收"}
POINT_DONE = {"已验收", "待验收"}
ACCEPT_DONE = {"已验收"}
BUG_FIXED = {"done", "closed", "converted"}
SEVERITY_W = {
    "Critical": 10.0,
    "Major": 5.0,
    "Medium": 2.0,
    "Minor": 1.0,
    "Trivial": 0.5,
}

HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "summary": ("summary",),
    "status": ("status",),
    "delay_reason": ("延迟原因",),
    "risk_report_reason": ("未提前上报风险的原因", "未提前上报风险原因"),
    "priority": ("priority",),
    "sprint": ("sprint",),
    "dev_owner": ("dev owner",),
    "qc_owner": ("qc owner",),
    "product_owner": ("产品经理", "产品 owner", "pm owner"),
    "task_owner": ("task owner",),
    "tech_owner": ("tech owner",),
    "story_point_dev": ("story point (dev)", "story points (dev)", "开发 故事", "开发故事"),
    "story_point_qc": ("story point (qc)", "story points (qc)"),
    "story_points": ("story points",),
    "url": ("user story链接", "task链接", "tech improvement链接", "bug链接", "链接"),
    "severity": ("severity",),
    "root_cause": ("root cause category",),
    "environment": ("bug environment", "environment"),
    "assignee": ("assignee",),
    "role": ("角色",),
    "estimated": ("estimated",),
    "parent_summary": ("关联工作项-标题", "关联工作项标题"),
    "parent_type": ("关联工作项-工作项类型", "关联工作项类型", "工作项类型"),
    "subtask_name": ("sub-task name", "subtask name"),
}


def _norm_header(v: Any) -> str:
    return re.sub(r"\s+", " ", str(v or "").strip().lower())


def _header_map(row: tuple[Any, ...]) -> dict[str, int]:
    out: dict[str, int] = {}
    for i, cell in enumerate(row):
        h = _norm_header(cell)
        if not h:
            continue
        for key, aliases in HEADER_ALIASES.items():
            if key in out:
                continue
            if h in aliases or any(a in h for a in aliases):
                out[key] = i
    return out


def _cell(row: tuple[Any, ...], hm: dict[str, int], key: str, default: Any = "") -> Any:
    idx = hm.get(key)
    if idx is None or idx >= len(row):
        return default
    v = row[idx]
    return default if v is None else v


def _float(v: Any) -> float:
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def _id_from_url(url: str) -> str:
    m = re.search(r"/(\d{8,})(?:\?|$)", str(url or ""))
    return m.group(1) if m else ""


def _is_sprint_banner(summary: str) -> bool:
    s = str(summary or "")
    return "共" in s and "项" in s and "OBIS-" in s


def _read_sheet(path: Path) -> tuple[dict[str, int], list[tuple[Any, ...]]]:
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    if not rows:
        return {}, []
    hm = _header_map(rows[0])
    return hm, rows[1:]


def _subtask_row(row: tuple[Any, ...], hm: dict[str, int]) -> dict[str, Any]:
    """Feishu subtask export: parent cols often at fixed indices 8/9/11/12."""
    name = str(_cell(row, hm, "subtask_name", row[0] if row else "")).strip()
    parent_summary = str(_cell(row, hm, "parent_summary", row[8] if len(row) > 8 else "")).strip()
    parent_type = str(_cell(row, hm, "parent_type", row[9] if len(row) > 9 else "")).strip()
    role = str(_cell(row, hm, "role", row[11] if len(row) > 11 else "")).strip().upper()
    est_raw = _cell(row, hm, "estimated", None)
    if est_raw in (None, "") and len(row) > 12:
        est_raw = row[12]
    if est_raw in (None, "") and len(row) > 6:
        est_raw = row[6]
    return {
        "name": name,
        "status": str(_cell(row, hm, "status", "")).strip(),
        "assignee": str(_cell(row, hm, "assignee", "")).strip(),
        "role": role,
        "estimated": _float(est_raw),
        "parentType": parent_type,
        "parentSummary": parent_summary,
    }


def _parse_parent_items(path: Path, item_type: str) -> list[dict[str, Any]]:
    hm, rows = _read_sheet(path)
    out: list[dict[str, Any]] = []
    for row in rows:
        summary = str(_cell(row, hm, "summary", "")).strip()
        status = str(_cell(row, hm, "status", "")).strip()
        if not summary or _is_banner_row(summary, status):
            continue
        if not status:
            continue
        url = str(_cell(row, hm, "url", "")).strip()
        status = str(_cell(row, hm, "status", "")).strip()
        dev = _float(_cell(row, hm, "story_point_dev", row[10] if len(row) > 10 else 0))
        qc = _float(_cell(row, hm, "story_point_qc", row[11] if len(row) > 11 else 0))
        sp = _float(_cell(row, hm, "story_points", row[10] if len(row) > 10 else 0))
        if item_type == "Task":
            dev, qc = sp, 0.0
        owner = (
            _cell(row, hm, "dev_owner")
            or _cell(row, hm, "task_owner")
            or _cell(row, hm, "tech_owner")
            or ""
        )
        rec: dict[str, Any] = {
            "type": item_type,
            "id": _id_from_url(url),
            "summary": summary,
            "status": status,
            "sprint": str(_cell(row, hm, "sprint", "")).strip(),
            "url": url,
            "owner": str(owner or "—").strip() or "—",
            "productOwner": str(_cell(row, hm, "product_owner", "")).strip(),
            "devOwner": str(owner or "").strip(),
            "qcOwner": str(_cell(row, hm, "qc_owner", "")).strip(),
            "pointDev": dev,
            "pointQc": qc,
            "pointTotal": round(dev + qc, 2),
            "delayReason": str(_cell(row, hm, "delay_reason", "")).strip(),
            "riskReportReason": str(_cell(row, hm, "risk_report_reason", "")).strip(),
        }
        out.append(rec)
    return out


def _parse_bugs(path: Path, sprint: str = "") -> list[dict[str, Any]]:
    hm, rows = _read_sheet(path)
    out: list[dict[str, Any]] = []
    for row in rows:
        summary = str(_cell(row, hm, "summary", "")).strip()
        status = str(_cell(row, hm, "status", "")).strip()
        if not summary or _is_banner_row(summary, status) or not status:
            continue
        sev = str(_cell(row, hm, "severity", "")).strip()
        if not sev:
            continue
        sp = str(_cell(row, hm, "sprint", "")).strip()
        if sprint and sp and sp != sprint:
            continue
        url = str(_cell(row, hm, "url", "")).strip()
        root = str(_cell(row, hm, "root_cause", "")).strip() or "TBD"
        env = str(_cell(row, hm, "environment", "")).strip()
        out.append(
            {
                "id": _id_from_url(url),
                "summary": summary,
                "status": status,
                "severity": sev,
                "rootCause": root,
                "environment": env,
                "url": url,
                "fixed": status.lower() in BUG_FIXED,
                "sprint": sp,
            }
        )
    return out


def _is_banner_row(summary: str, status: str = "") -> bool:
    s = str(summary or "")
    if _is_sprint_banner(s):
        return True
    if s and not str(status or "").strip():
        if "共" in s and "条" in s:
            return True
    return False


def _is_prod_env(raw: str) -> bool:
    key = (raw or "").strip().upper()
    return key in {"PRD", "PROD", "PRODUCTION", "生产", "生产环境"}


def _parse_subtasks(path: Path) -> list[dict[str, Any]]:
    hm, rows = _read_sheet(path)
    out: list[dict[str, Any]] = []
    for row in rows:
        rec = _subtask_row(row, hm)
        if not rec["name"]:
            continue
        out.append(rec)
    return out


def _pct(n: float, d: float) -> float | None:
    if not d:
        return None
    return round(n * 100.0 / d, 1)


def _pct_str(n: float, d: float) -> str:
    p = _pct(n, d)
    return f"{p}%" if p is not None else "—"


def _type_delivery(rows: list[dict[str, Any]], *, item_done: set[str]) -> dict[str, Any]:
    total = len(rows)
    done = sum(1 for r in rows if r.get("status") in item_done)
    unfin = total - done
    rate = _pct(done, total)
    return {
        "total": total,
        "done": done,
        "unfinished": unfin,
        "doneRate": rate,
        "doneRateLabel": _pct_str(done, total),
    }


def _role_points(
    subtasks: list[dict[str, Any]],
    parents: dict[str, list[dict[str, Any]]],
    parent_done: set[str],
) -> dict[str, Any]:
    parent_index: dict[tuple[str, str], dict[str, Any]] = {}
    for typ, rows in parents.items():
        for r in rows:
            parent_index[(typ, r["summary"])] = r

    plan_by_type: dict[str, dict[str, float]] = {}
    done_by_type: dict[str, dict[str, float]] = {}

    for st in subtasks:
        pt = st.get("parentType") or ""
        ps = st.get("parentSummary") or ""
        parent = parent_index.get((pt, ps))
        if not parent:
            continue
        role = st.get("role") or ""
        if role not in ("DEV", "QC"):
            continue
        est = _float(st.get("estimated"))
        plan = plan_by_type.setdefault(pt, {"planDev": 0.0, "planQc": 0.0})
        if role == "DEV":
            plan["planDev"] += est
        else:
            plan["planQc"] += est
        if parent.get("status") in parent_done:
            done = done_by_type.setdefault(pt, {"doneDev": 0.0, "doneQc": 0.0})
            if role == "DEV":
                done["doneDev"] += est
            else:
                done["doneQc"] += est

    out: dict[str, Any] = {}
    for typ in ("User Story", "Task", "Tech Improvement"):
        plan = plan_by_type.get(typ, {"planDev": 0.0, "planQc": 0.0})
        done = done_by_type.get(typ, {"doneDev": 0.0, "doneQc": 0.0})
        pd, pq = plan["planDev"], plan["planQc"]
        dd, dq = done.get("doneDev", 0.0), done.get("doneQc", 0.0)
        out[typ] = {
            "planDev": round(pd, 1),
            "planQc": round(pq, 1),
            "doneDev": round(dd, 1),
            "doneQc": round(dq, 1),
            "devRate": _pct(dd, pd),
            "qcRate": _pct(dq, pq),
        }
    # totals
    pd = sum(out[t]["planDev"] for t in out)
    pq = sum(out[t]["planQc"] for t in out)
    dd = sum(out[t]["doneDev"] for t in out)
    dq = sum(out[t]["doneQc"] for t in out)
    out["total"] = {
        "planDev": round(pd, 1),
        "planQc": round(pq, 1),
        "doneDev": round(dd, 1),
        "doneQc": round(dq, 1),
        "devRate": _pct(dd, pd),
        "qcRate": _pct(dq, pq),
        "planTotal": round(pd + pq, 1),
        "doneTotal": round(dd + dq, 1),
        "totalRate": _pct(dd + dq, pd + pq),
    }
    return out


def parse_sprint_xlsx_bundle(
    *,
    user_story: Path,
    tech_improvement: Path,
    task: Path,
    bug: Path,
    subtasks: Path,
) -> dict[str, Any]:
    stories = _parse_parent_items(user_story, "User Story")
    tasks = _parse_parent_items(task, "Task")
    tis = _parse_parent_items(tech_improvement, "Tech Improvement")
    subtask_rows = _parse_subtasks(subtasks)

    sprint = ""
    for r in stories + tasks + tis:
        sp = str(r.get("sprint") or "").strip()
        if sp:
            sprint = sp
            break
    if not sprint:
        _, rows = _read_sheet(user_story)
        for row in rows:
            m = re.search(r"(OBIS-\d{8}-\d{8})", str(row[0] if row else ""))
            if m:
                sprint = m.group(1)
                break

    bugs = _parse_bugs(bug, sprint)

    parents = {
        "User Story": stories,
        "Task": tasks,
        "Tech Improvement": tis,
    }
    role_pts = _role_points(subtask_rows, parents, POINT_DONE)

    all_items = stories + tasks + tis
    accept_done = sum(1 for r in stories if r.get("status") in ACCEPT_DONE)
    item_total = len(stories)

    sev_counts: dict[str, dict[str, int]] = {}
    base_dqs = 0.0
    for b in bugs:
        sev = b.get("severity") or "(空)"
        sev_counts.setdefault(sev, {"found": 0, "fixed": 0})
        sev_counts[sev]["found"] += 1
        if b.get("fixed"):
            sev_counts[sev]["fixed"] += 1
        base_dqs += SEVERITY_W.get(sev, 0.0)

    bug_fixed = sum(1 for b in bugs if b.get("fixed"))
    bug_total = len(bugs)

    root_counts: dict[str, int] = {}
    for b in bugs:
        rc = b.get("rootCause") or "TBD"
        if not str(rc).strip():
            rc = "TBD"
        root_counts[rc] = root_counts.get(rc, 0) + 1

    # 未完成清单仅保留 User Story（报告第 6 节复盘表不含 Task / Tech Improvement）
    incomplete = [
        r
        for r in all_items
        if r.get("status") not in ITEM_DONE and r.get("type") == "User Story"
    ]

    dev_people: set[str] = set()
    qc_people: set[str] = set()
    for st in subtask_rows:
        a = str(st.get("assignee") or "").strip()
        if not a:
            continue
        role = str(st.get("role") or "").strip().upper()
        if role == "QC":
            qc_people.add(a)
        else:
            dev_people.add(a)
    people_by_role = {
        "DEV": len(dev_people),
        "QC": len(qc_people),
        "PM": None,
    }

    dev_done_pts = role_pts["total"]["doneDev"]
    norm_dqs = round(base_dqs / dev_done_pts, 2) if dev_done_pts else None

    return {
        "sprint": sprint,
        "source": "飞书 Sprint 导出 xlsx",
        "stories": stories,
        "tasks": tasks,
        "techImprovements": tis,
        "bugs": bugs,
        "subtasks": subtask_rows,
        "metrics": {
            "acceptClosed": accept_done,
            "acceptTotal": item_total,
            "acceptRate": _pct(accept_done, item_total),
            "acceptLabel": f"{accept_done}/{item_total}",
            "pointDone": role_pts["total"]["doneTotal"],
            "pointPlan": role_pts["total"]["planTotal"],
            "pointRate": role_pts["total"]["totalRate"],
            "bugFixed": bug_fixed,
            "bugTotal": bug_total,
            "bugFixRate": _pct(bug_fixed, bug_total),
            "baseDqs": round(base_dqs, 1),
            "normDqs": norm_dqs,
        },
        "delivery": {
            "User Story": _type_delivery(stories, item_done=ITEM_DONE),
            "Task": _type_delivery(tasks, item_done={"待验收", "已验收"}),
            "Tech Improvement": _type_delivery(tis, item_done=ITEM_DONE),
        },
        "rolePoints": role_pts,
        "severity": sev_counts,
        "rootCause": root_counts,
        "incompleteItems": incomplete,
        "peopleCount": len(dev_people | qc_people),
        "peopleByRole": people_by_role,
        "prodBugCount": sum(
            1 for b in bugs if _is_prod_env(str(b.get("environment") or ""))
        ),
    }


def default_xlsx_paths(base_dir: Path) -> dict[str, Path]:
    """Resolve standard export filenames under a directory."""
    mapping = {
        "user_story": "User Story导出-OBIS (31).xlsx",
        "tech_improvement": "Tech Improvement导出-OBIS (13).xlsx",
        "task": "Task导出-OBIS (21).xlsx",
        "bug": "Bug导出-OBIS (20).xlsx",
        "subtasks": "任务导出-OBIS (36).xlsx",
    }
    out: dict[str, Path] = {}
    for key, name in mapping.items():
        p = base_dir / name
        if not p.exists():
            # fuzzy match by prefix
            hits = list(base_dir.glob(name.split("导出")[0] + "*.xlsx"))
            if hits:
                p = hits[0]
        out[key] = p
    return out


def parse_default_bundle(base_dir: Path) -> dict[str, Any]:
    paths = default_xlsx_paths(base_dir)
    missing = [k for k, p in paths.items() if not p.exists()]
    if missing:
        raise FileNotFoundError(f"缺少导出文件: {', '.join(missing)} @ {base_dir}")
    return parse_sprint_xlsx_bundle(
        user_story=paths["user_story"],
        tech_improvement=paths["tech_improvement"],
        task=paths["task"],
        bug=paths["bug"],
        subtasks=paths["subtasks"],
    )
